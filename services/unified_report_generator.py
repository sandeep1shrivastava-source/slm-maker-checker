import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

def build_master_audit_excel(academic_results: dict, prompt_results: dict = None, metadata: dict = None) -> bytes:
    wb = Workbook()

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))

    def apply_sheet_styling(ws, df):
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_col=ws.max_column, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.border = border
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(len(str(col)) + 5, 22)

    # Sheet 1: Master Executive Dashboard
    ws_dash = wb.active
    ws_dash.title = "Master Executive Dashboard"

    uni = metadata.get("university_name", "N/A") if metadata else "N/A"
    prog = metadata.get("program_name", "N/A") if metadata else "N/A"
    deliv = metadata.get("deliverable_type", "SLM Unit") if metadata else "SLM Unit"
    user_email = metadata.get("operator_email", "Unassigned") if metadata else "Unassigned"
    user_role = metadata.get("operator_role", "Lead Administrator") if metadata else "Lead Administrator"

    p_score = f"{prompt_results.get('prompt_score', 0):.1f}%" if prompt_results else "Stage 3 Pending"
    p_status = prompt_results.get("overall_compliance_status", "NOT RUN") if prompt_results else "NOT RUN"

    dash_metrics = [
        ["Audit Parameter / Pipeline Layer", "Recorded Value", "Target Benchmark"],
        ["Target University", uni, "Accredited University"],
        ["Academic Program / Discipline", prog, "Approved Curriculum"],
        ["Deliverable Type", deliv, "Mandatory Scope"],
        ["Assigned Reviewer Email", user_email, "Authorized Sign-off"],
        ["Reviewer Role & Permission", user_role, "Role-Based Access"],
        ["Stage 2: Academic Gate Decision", academic_results.get("decision", "REVISE"), "PASS"],
        ["Stage 2: Quality Score", f"{academic_results.get('score', 0):.1f}%", ">= 95.0%"],
        ["Stage 2: Syllabus Coverage", f"{academic_results.get('coverage_pct', 0):.1f}%", "100.0%"],
        ["Stage 2: Critical Errors", str(academic_results.get("critical_count", 0)), "0"],
        ["Stage 3: Prompt Compliance Score", p_score, ">= 95.0%"],
        ["Stage 3: Prompt Compliance Status", p_status, "COMPLIANT"]
    ]
    for row in dash_metrics:
        ws_dash.append(row)
    for cell in ws_dash[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    for row in ws_dash.iter_rows(min_row=2, max_col=3, max_row=len(dash_metrics)):
        for cell in row:
            cell.border = border
    ws_dash.column_dimensions['A'].width = 34
    ws_dash.column_dimensions['B'].width = 38
    ws_dash.column_dimensions['C'].width = 25

    # Sheet 2: Academic Error Register
    ws_err = wb.create_sheet(title="Academic Error Register")
    df_err = pd.DataFrame(academic_results.get("error_register", []))
    if not df_err.empty:
        apply_sheet_styling(ws_err, df_err)

    # Sheet 3: Syllabus Matrix
    ws_syl = wb.create_sheet(title="Syllabus Matrix")
    df_syl = pd.DataFrame(academic_results.get("syllabus_mapping", []))
    if not df_syl.empty:
        apply_sheet_styling(ws_syl, df_syl)

    # Sheet 4: Prompt Adherence Checks
    if prompt_results and "compliance_checks" in prompt_results:
        ws_prompt = wb.create_sheet(title="Prompt Adherence Checks")
        df_p = pd.DataFrame(prompt_results["compliance_checks"])
        if not df_p.empty:
            apply_sheet_styling(ws_prompt, df_p)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()