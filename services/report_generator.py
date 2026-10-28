import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

def build_audit_excel(audit_results: dict) -> bytes:
    wb = Workbook()

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))

    def apply_sheet_styling(ws, df):
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_col=ws.max_column, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.border = border
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(len(str(col)) + 5, 22)

    # 1. Dashboard Tab
    ws_dash = wb.active
    ws_dash.title = "Audit Dashboard"
    dash_metrics = [
        ["Metric", "Value", "Benchmark Threshold"],
        ["Audit Decision", audit_results.get("decision", "REVISE"), "PASS"],
        ["Weighted Quality Score", f"{audit_results.get('score', 0):.1f}%", ">= 95.0%"],
        ["Syllabus Coverage", f"{audit_results.get('coverage_pct', 0):.1f}%", "100.0%"],
        ["Prompt Adherence Status", audit_results.get("prompt_status", "FAIL"), "PASS"],
        ["Critical Errors Unresolved", str(audit_results.get("critical_count", 0)), "0"]
    ]
    for row in dash_metrics:
        ws_dash.append(row)
    for cell in ws_dash[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    ws_dash.column_dimensions['A'].width = 30
    ws_dash.column_dimensions['B'].width = 30
    ws_dash.column_dimensions['C'].width = 25

    # 2. AI Prompt Adherence Tab
    ws_prompt = wb.create_sheet(title="AI Prompt Adherence")
    df_prompt = pd.DataFrame(audit_results.get("prompt_adherence_logs", []))
    if not df_prompt.empty:
        apply_sheet_styling(ws_prompt, df_prompt)

    # 3. Syllabus Mapping Tab
    ws_syllabus = wb.create_sheet(title="Syllabus Mapping")
    df_syl = pd.DataFrame(audit_results.get("syllabus_mapping", []))
    if not df_syl.empty:
        apply_sheet_styling(ws_syllabus, df_syl)

    # 4. Error Register Tab
    ws_err = wb.create_sheet(title="Error Register")
    df_err = pd.DataFrame(audit_results.get("error_register", []))
    if not df_err.empty:
        apply_sheet_styling(ws_err, df_err)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()