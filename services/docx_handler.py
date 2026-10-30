import io
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def extract_text_from_docx(file_input) -> str:
    """Extracts text content safely from uploaded docx files, raw bytes, or file buffers."""
    if file_input is None:
        return ""

    # Handle Streamlit UploadedFile buffers or raw bytes
    if hasattr(file_input, "getvalue"):
        file_bytes = file_input.getvalue()
    elif hasattr(file_input, "read"):
        file_bytes = file_input.read()
    else:
        file_bytes = file_input

    if isinstance(file_bytes, bytes):
        doc = Document(io.BytesIO(file_bytes))
    else:
        doc = Document(file_input)

    full_text = []
    for p in doc.paragraphs:
        if p.text.strip():
            full_text.append(p.text.strip())
            
    for tbl in doc.tables:
        for row in tbl.rows:
            row_vals = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_vals:
                full_text.append(" | ".join(row_vals))

    return "\n".join(full_text)

def generate_slm_docx(title: str, content: str) -> io.BytesIO:
    """Generates a downloadable DOCX for the generated academic deliverable."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph("―" * 50)
    for block in content.split("\n\n"):
        if block.strip():
            doc.add_paragraph(block.strip())

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream

def generate_prompt_audit_docx(audit_data: dict, deliverable_type: str = "SLM", university: str = "State Open University", program: str = "MBA") -> io.BytesIO:
    """Generates an executive-grade Word document audit dossier for Stage 3 Prompt Compliance."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("ACADEMIC PROMPT & INSTRUCTION AUDIT REPORT")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0, 51, 102)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Stage 3 Deep Meta-Audit Dossier | Deliverable: {deliverable_type}")
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("―" * 55)

    meta_tbl = doc.add_table(rows=2, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.rows[0].cells[0].text = f"University: {university}"
    meta_tbl.rows[0].cells[1].text = f"Audit Status: {audit_data.get('overall_compliance_status', 'COMPLIANT')}"
    meta_tbl.rows[1].cells[0].text = f"Academic Program: {program}"
    meta_tbl.rows[1].cells[1].text = f"Prompt Score: {audit_data.get('prompt_score', 0.0)} / 100"

    doc.add_paragraph("")

    h1 = doc.add_heading("1. Executive Compliance Summary", level=2)
    h1.style.font.color.rgb = RGBColor(0, 51, 102)

    p_sum = doc.add_paragraph()
    p_sum.add_run("• Compliance Status: ").bold = True
    p_sum.add_run(f"{audit_data.get('overall_compliance_status', 'COMPLIANT')}\n")
    p_sum.add_run("• Prompt Fidelity Score: ").bold = True
    p_sum.add_run(f"{audit_data.get('prompt_score', 0.0)}%\n")
    p_sum.add_run("• Negative Constraint Violations: ").bold = True
    p_sum.add_run(f"{audit_data.get('negative_constraint_violations', 0)} detected\n")

    h2 = doc.add_heading("2. Pedagogical & Prompt Compliance Matrix", level=2)
    h2.style.font.color.rgb = RGBColor(0, 51, 102)

    checks = audit_data.get("compliance_checks", [])
    if checks:
        tbl = doc.add_table(rows=1, cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.style = "Table Grid"

        headers = ["Dimension", "Status", "Severity", "Auditor Findings"]
        for idx, text in enumerate(headers):
            cell = tbl.rows[0].cells[idx]
            cell.text = text
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True

        for item in checks:
            row = tbl.add_row().cells
            row[0].text = str(item.get("dimension", ""))
            row[1].text = str(item.get("status", ""))
            row[2].text = str(item.get("severity", "None"))
            row[3].text = str(item.get("finding", ""))

    doc.add_paragraph("")

    h3 = doc.add_heading("3. Structural Drift & Guardrail Analysis", level=2)
    h3.style.font.color.rgb = RGBColor(0, 51, 102)
    doc.add_paragraph(audit_data.get("structural_drift_analysis", "Zero structural drift detected."))

    doc.add_paragraph("\n" + "―" * 55)
    footer = doc.add_paragraph("Certified by University Academic AI Meta-Auditor")
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style.font.italic = True
    footer.style.font.size = Pt(9)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out